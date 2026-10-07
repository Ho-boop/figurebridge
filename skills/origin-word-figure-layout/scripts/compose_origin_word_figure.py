from __future__ import annotations

import argparse
import copy
import json
import math
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from PIL import Image


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PIC = "http://schemas.openxmlformats.org/drawingml/2006/picture"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"w": W, "wp": WP, "a": A, "pic": PIC, "r": R, "ct": CT}
EMU_PER_INCH = 914400
SUPPORTED = {".emf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
CONTENT_TYPES = {
    ".emf": "image/x-emf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}


@dataclass
class Panel:
    source: Path
    preview: Path
    width: int
    height: int
    content_box: tuple[int, int, int, int]
    frame: tuple[int, int, int, int]
    mask: np.ndarray
    touches_source_edge: bool
    opaque_box: bool


@dataclass
class Placement:
    anchor_x: float
    anchor_y: float
    width: int
    height: int


def natural_key(path: Path) -> list[object]:
    return [int(token) if token.isdigit() else token.casefold() for token in re.split(r"(\d+)", path.name)]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Crop, align, and compact Origin-exported panels into a Word DOCX."
    )
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--cols", type=int, required=True)
    parser.add_argument("--files", nargs="*")
    parser.add_argument("--qa-dir", type=Path)
    parser.add_argument("--frame-overrides-json", type=Path)
    parser.add_argument("--max-width-in", type=float, default=5.75)
    parser.add_argument("--max-height-in", type=float, default=9.0)
    parser.add_argument("--safety-pt", type=float, default=1.0)
    parser.add_argument("--preview-width", type=int, default=2600)
    parser.add_argument("--page-width-in", type=float, default=8.27)
    parser.add_argument("--page-height-in", type=float, default=11.69)
    parser.add_argument("--left-margin-in", type=float, default=1.25)
    parser.add_argument("--right-margin-in", type=float, default=1.25)
    parser.add_argument("--top-margin-in", type=float, default=1.0)
    parser.add_argument("--bottom-margin-in", type=float, default=1.0)
    parser.add_argument("--top-in", type=float, default=1.0)
    parser.add_argument("--allow-edge-content", action="store_true")
    return parser.parse_args()


def resolve_sources(args: argparse.Namespace) -> list[Path]:
    folder = args.input_dir.resolve()
    if not folder.is_dir():
        raise FileNotFoundError(f"Input directory does not exist: {folder}")
    if args.files:
        sources = []
        for value in args.files:
            path = Path(value)
            if not path.is_absolute():
                path = folder / path
            sources.append(path.resolve())
    else:
        sources = sorted(
            (path.resolve() for path in folder.iterdir() if path.suffix.lower() in SUPPORTED),
            key=natural_key,
        )
    if not sources:
        raise RuntimeError("No supported panel files were found.")
    for source in sources:
        if not source.is_file():
            raise FileNotFoundError(source)
        if source.suffix.lower() not in SUPPORTED:
            raise RuntimeError(f"Unsupported panel format: {source}")
    if len(sources) % args.cols:
        raise RuntimeError(
            f"{len(sources)} panels do not fill a rectangular grid with {args.cols} columns."
        )
    kinds = {"emf" if path.suffix.lower() == ".emf" else "raster" for path in sources}
    if len(kinds) != 1:
        raise RuntimeError("Do not mix EMF and raster panels in one run; their source scales are ambiguous.")
    return sources


def render_previews(
    sources: list[Path], qa_dir: Path, preview_width: int
) -> list[Path]:
    preview_dir = qa_dir / "previews"
    preview_dir.mkdir(parents=True, exist_ok=True)
    previews = [preview_dir / f"panel-{index:03d}.png" for index in range(1, len(sources) + 1)]
    if sources[0].suffix.lower() != ".emf":
        for source, output in zip(sources, previews):
            with Image.open(source) as image:
                background = Image.new("RGB", image.size, "white")
                if image.mode in ("RGBA", "LA"):
                    alpha = image.getchannel("A")
                    background.paste(image.convert("RGB"), mask=alpha)
                else:
                    background.paste(image.convert("RGB"))
                target_height = max(1, round(background.height * preview_width / background.width))
                background.resize(
                    (preview_width, target_height), Image.Resampling.LANCZOS
                ).save(output, "PNG")
        return previews

    ascii_dir = qa_dir / "emf-analysis-sources"
    ascii_dir.mkdir(parents=True, exist_ok=True)
    manifest = []
    for index, (source, output) in enumerate(zip(sources, previews), start=1):
        ascii_source = ascii_dir / f"panel-{index:03d}.emf"
        shutil.copyfile(source, ascii_source)
        manifest.append({"source": str(ascii_source), "output": output.name})
    manifest_path = qa_dir / "emf-preview-manifest.json"
    # ASCII-escaped JSON remains reliable with Windows PowerShell 5.1's legacy
    # default text decoding while ConvertFrom-Json restores Unicode paths.
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=True, indent=2), encoding="ascii")

    powershell = shutil.which("powershell.exe") or str(
        Path("C:/Windows/System32/WindowsPowerShell/v1.0/powershell.exe")
    )
    renderer = Path(__file__).with_name("render_emf_previews.ps1")
    command = [
        powershell,
        "-NoProfile",
        "-ExecutionPolicy",
        "Bypass",
        "-File",
        str(renderer),
        "-Manifest",
        str(manifest_path),
        "-OutputDir",
        str(preview_dir),
        "-TargetWidth",
        str(preview_width),
    ]
    completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
    if completed.returncode:
        raise RuntimeError(
            "EMF preview rendering failed.\n"
            + completed.stdout
            + "\n"
            + completed.stderr
        )
    return previews


def contiguous_groups(values: np.ndarray) -> list[tuple[int, int]]:
    if values.size == 0:
        return []
    groups: list[tuple[int, int]] = []
    start = previous = int(values[0])
    for value in map(int, values[1:]):
        if value != previous + 1:
            groups.append((start, previous + 1))
            start = value
        previous = value
    groups.append((start, previous + 1))
    return groups


def longest_run(vector: np.ndarray) -> int:
    indices = np.flatnonzero(vector)
    if indices.size == 0:
        return 0
    breaks = np.flatnonzero(np.diff(indices) > 1)
    starts = np.r_[0, breaks + 1]
    ends = np.r_[breaks, indices.size - 1]
    return int(np.max(indices[ends] - indices[starts] + 1))


def line_candidates(mask: np.ndarray, horizontal: bool) -> list[int]:
    length = mask.shape[1] if horizontal else mask.shape[0]
    count = mask.shape[0] if horizontal else mask.shape[1]
    scores = np.empty(count, dtype=np.int32)
    for index in range(count):
        vector = mask[index, :] if horizontal else mask[:, index]
        scores[index] = longest_run(vector)
    threshold = max(20, round(length * 0.22))
    groups = contiguous_groups(np.flatnonzero(scores >= threshold))
    return [
        max(range(start, end), key=lambda value: int(scores[value]))
        for start, end in groups
    ]


def detect_frame(dark_mask: np.ndarray, source_name: str) -> tuple[int, int, int, int]:
    rows = line_candidates(dark_mask, horizontal=True)
    cols = line_candidates(dark_mask, horizontal=False)
    if len(rows) < 2 or len(cols) < 2:
        raise RuntimeError(
            f"Could not detect a rectangular plot frame in {source_name}. "
            "Supply --frame-overrides-json after inspecting the QA preview."
        )
    top, bottom = min(rows), max(rows)
    left, right = min(cols), max(cols)
    height, width = dark_mask.shape
    if right - left < width * 0.25 or bottom - top < height * 0.25:
        raise RuntimeError(
            f"Implausible plot-frame detection in {source_name}: "
            f"{left, top, right, bottom}. Supply --frame-overrides-json."
        )
    return left, top, right, bottom


def load_overrides(path: Path | None) -> dict[str, tuple[int, int, int, int]]:
    if path is None:
        return {}
    raw = json.loads(path.read_text(encoding="utf-8"))
    result = {}
    for name, value in raw.items():
        if isinstance(value, dict):
            box = (value["left"], value["top"], value["right"], value["bottom"])
        else:
            box = tuple(value)
        if len(box) != 4:
            raise ValueError(f"Frame override for {name} must have four values.")
        result[name] = tuple(map(int, box))
    return result


def analyze_panel(source: Path, preview: Path, override: tuple[int, int, int, int] | None) -> Panel:
    with Image.open(preview) as image:
        rgb = np.asarray(image.convert("RGB"))
    height, width = rgb.shape[:2]
    spread = rgb.max(axis=2).astype(np.int16) - rgb.min(axis=2).astype(np.int16)
    mask = (rgb.min(axis=2) < 246) | (spread > 5)
    ys, xs = np.where(mask)
    if xs.size == 0:
        raise RuntimeError(f"Panel appears blank: {source}")
    touches = bool(
        mask[0, :].any() or mask[-1, :].any() or mask[:, 0].any() or mask[:, -1].any()
    )
    padding = max(10, round(min(width, height) * 0.004))
    x0 = max(0, int(xs.min()) - padding)
    y0 = max(0, int(ys.min()) - padding)
    x1 = min(width, int(xs.max()) + 1 + padding)
    y1 = min(height, int(ys.max()) + 1 + padding)
    dark = rgb.max(axis=2) < 110
    frame = override or detect_frame(dark, source.name)
    left, top, right, bottom = frame
    if not (0 <= left < right < width and 0 <= top < bottom < height):
        raise RuntimeError(f"Frame is outside the preview for {source.name}: {frame}")
    return Panel(
        source,
        preview,
        width,
        height,
        (x0, y0, x1, y1),
        frame,
        mask[y0:y1, x0:x1],
        touches,
        source.suffix.lower() != ".emf",
    )


def row_profile(panel: Panel) -> dict[int, tuple[int, int]]:
    mask = panel.mask
    x0, y0, _, _ = panel.content_box
    left, top, _, _ = panel.frame
    result = {}
    for local_y in np.flatnonzero(mask.any(axis=1)):
        xs = np.flatnonzero(mask[local_y])
        relative_y = int(local_y + y0 - top)
        result[relative_y] = (int(xs[0] + x0 - left), int(xs[-1] + x0 - left))
    return result


def column_profile(panel: Panel) -> dict[int, tuple[int, int]]:
    mask = panel.mask
    x0, y0, _, _ = panel.content_box
    left, top, _, _ = panel.frame
    result = {}
    for local_x in np.flatnonzero(mask.any(axis=0)):
        ys = np.flatnonzero(mask[:, local_x])
        relative_x = int(local_x + x0 - left)
        result[relative_x] = (int(ys[0] + y0 - top), int(ys[-1] + y0 - top))
    return result


def horizontal_requirement(
    left_profile: dict[int, tuple[int, int]],
    right_profile: dict[int, tuple[int, int]],
    clearance: int,
) -> int:
    common = left_profile.keys() & right_profile.keys()
    if not common:
        return 1
    return max(left_profile[y][1] - right_profile[y][0] + 1 + clearance for y in common)


def vertical_requirement(
    top_profile: dict[int, tuple[int, int]],
    bottom_profile: dict[int, tuple[int, int]],
    clearance: int,
) -> int:
    common = top_profile.keys() & bottom_profile.keys()
    if not common:
        return 1
    return max(top_profile[x][1] - bottom_profile[x][0] + 1 + clearance for x in common)


def solve_layout(
    panels: list[Panel], rows: int, cols: int, clearance: int
) -> tuple[list[Placement], float, float, list[float], list[float]]:
    row_profiles = [row_profile(panel) for panel in panels]
    column_profiles = [column_profile(panel) for panel in panels]
    horizontal: dict[tuple[int, int], int] = {}
    for first in range(cols):
        for second in range(first + 1, cols):
            requirements = []
            for row in range(rows):
                left_panel = panels[row * cols + first]
                right_panel = panels[row * cols + second]
                requirement = horizontal_requirement(
                    row_profiles[row * cols + first],
                    row_profiles[row * cols + second],
                    clearance,
                )
                if left_panel.opaque_box or right_panel.opaque_box:
                    _, _, left_x1, _ = left_panel.content_box
                    right_x0, _, _, _ = right_panel.content_box
                    left_frame_x = left_panel.frame[0]
                    right_frame_x = right_panel.frame[0]
                    requirement = max(
                        requirement,
                        (left_x1 - left_frame_x) - (right_x0 - right_frame_x),
                    )
                requirements.append(requirement)
            horizontal[first, second] = max(requirements)
    frame_x = [0.0] * cols
    for second in range(1, cols):
        frame_x[second] = max(
            frame_x[first] + horizontal[first, second] for first in range(second)
        )

    vertical: dict[tuple[int, int], int] = {}
    for first in range(rows):
        for second in range(first + 1, rows):
            requirements = []
            for col in range(cols):
                top_panel = panels[first * cols + col]
                bottom_panel = panels[second * cols + col]
                requirement = vertical_requirement(
                    column_profiles[first * cols + col],
                    column_profiles[second * cols + col],
                    clearance,
                )
                if top_panel.opaque_box or bottom_panel.opaque_box:
                    _, _, _, top_y1 = top_panel.content_box
                    _, bottom_y0, _, _ = bottom_panel.content_box
                    top_frame_y = top_panel.frame[1]
                    bottom_frame_y = bottom_panel.frame[1]
                    requirement = max(
                        requirement,
                        (top_y1 - top_frame_y) - (bottom_y0 - bottom_frame_y),
                    )
                requirements.append(requirement)
            vertical[first, second] = max(requirements)
    frame_y = [0.0] * rows
    for second in range(1, rows):
        frame_y[second] = max(
            frame_y[first] + vertical[first, second] for first in range(second)
        )

    placements = []
    for index, panel in enumerate(panels):
        row, col = divmod(index, cols)
        x0, y0, x1, y1 = panel.content_box
        left, top, _, _ = panel.frame
        placements.append(
            Placement(
                frame_x[col] + x0 - left,
                frame_y[row] + y0 - top,
                x1 - x0,
                y1 - y0,
            )
        )
    minimum_x = min(item.anchor_x for item in placements)
    minimum_y = min(item.anchor_y for item in placements)
    for item in placements:
        item.anchor_x -= minimum_x
        item.anchor_y -= minimum_y
    width = max(item.anchor_x + item.width for item in placements)
    height = max(item.anchor_y + item.height for item in placements)
    return placements, width, height, frame_x, frame_y


def choose_layout(
    panels: list[Panel], rows: int, cols: int, max_width: float, max_height: float, safety_pt: float
) -> tuple[list[Placement], float, float, float, int, list[float], list[float]]:
    clearance = 0
    for _ in range(6):
        placements, width_px, height_px, frame_x, frame_y = solve_layout(
            panels, rows, cols, clearance
        )
        scale = min(max_width / width_px, max_height / height_px)
        next_clearance = max(0, math.ceil((safety_pt / 72.0) / scale))
        if next_clearance == clearance:
            break
        clearance = next_clearance
    placements, width_px, height_px, frame_x, frame_y = solve_layout(
        panels, rows, cols, clearance
    )
    scale = min(max_width / width_px, max_height / height_px)
    return placements, width_px, height_px, scale, clearance, frame_x, frame_y


def create_placeholder(path: Path, index: int) -> None:
    image = Image.new("RGB", (3, 3), (255, 255, 255))
    image.putpixel((1, 1), (240 + index % 15, 250 - index % 10, 245))
    image.save(path, "PNG")


def configure_document(args: argparse.Namespace, count: int, temp_dir: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(args.page_width_in)
    section.page_height = Inches(args.page_height_in)
    section.left_margin = Inches(args.left_margin_in)
    section.right_margin = Inches(args.right_margin_in)
    section.top_margin = Inches(args.top_margin_in)
    section.bottom_margin = Inches(args.bottom_margin_in)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    for index in range(1, count + 1):
        placeholder = temp_dir / f"placeholder-{index:03d}.png"
        create_placeholder(placeholder, index)
        paragraph.add_run().add_picture(str(placeholder), width=Inches(1))
    path = temp_dir / "base.docx"
    doc.save(path)
    return path


def wp_child(parent: etree._Element, name: str, **attrs: str) -> etree._Element:
    element = etree.SubElement(parent, f"{{{WP}}}{name}")
    for key, value in attrs.items():
        element.set(key, value)
    return element


def inline_to_anchor(inline: etree._Element, x_emu: int, y_emu: int, index: int) -> None:
    anchor = etree.Element(
        f"{{{WP}}}anchor",
        nsmap=inline.nsmap,
        distT="0",
        distB="0",
        distL="0",
        distR="0",
        simplePos="0",
        relativeHeight=str(251658240 + index),
        behindDoc="0",
        locked="0",
        layoutInCell="0",
        allowOverlap="1",
    )
    wp_child(anchor, "simplePos", x="0", y="0")
    position_h = wp_child(anchor, "positionH", relativeFrom="page")
    wp_child(position_h, "posOffset").text = str(x_emu)
    position_v = wp_child(anchor, "positionV", relativeFrom="page")
    wp_child(position_v, "posOffset").text = str(y_emu)
    extent = inline.find(f"{{{WP}}}extent")
    if extent is None:
        raise RuntimeError("An inline placeholder has no extent.")
    anchor.append(copy.deepcopy(extent))
    effect = inline.find(f"{{{WP}}}effectExtent")
    if effect is None:
        wp_child(anchor, "effectExtent", l="0", t="0", r="0", b="0")
    else:
        anchor.append(copy.deepcopy(effect))
    wp_child(anchor, "wrapNone")
    for name in ("docPr", "cNvGraphicFramePr"):
        node = inline.find(f"{{{WP}}}{name}")
        if node is not None:
            anchor.append(copy.deepcopy(node))
    graphic = inline.find(f"{{{A}}}graphic")
    if graphic is None:
        raise RuntimeError("An inline placeholder has no DrawingML graphic.")
    anchor.append(copy.deepcopy(graphic))
    inline.getparent().replace(inline, anchor)


def ensure_content_type(tree: etree._ElementTree, extension: str) -> None:
    extension = extension.lstrip(".").lower()
    existing = tree.xpath(
        f"/ct:Types/ct:Default[translate(@Extension, 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz')='{extension}']",
        namespaces=NS,
    )
    if existing:
        existing[0].set("ContentType", CONTENT_TYPES["." + extension])
        return
    node = etree.SubElement(tree.getroot(), f"{{{CT}}}Default")
    node.set("Extension", extension)
    node.set("ContentType", CONTENT_TYPES["." + extension])


def patch_docx(
    base: Path,
    output: Path,
    panels: list[Panel],
    placements: list[Placement],
    scale: float,
    group_left: float,
    group_top: float,
) -> None:
    with tempfile.TemporaryDirectory(prefix="origin_word_figure_patch_") as temp:
        root = Path(temp)
        with zipfile.ZipFile(base) as archive:
            archive.extractall(root)
        parser = etree.XMLParser(remove_blank_text=False)
        document_path = root / "word" / "document.xml"
        rels_path = root / "word" / "_rels" / "document.xml.rels"
        content_types_path = root / "[Content_Types].xml"
        document = etree.parse(str(document_path), parser)
        rels = etree.parse(str(rels_path), parser)
        content_types = etree.parse(str(content_types_path), parser)
        inlines = document.xpath("//wp:inline", namespaces=NS)
        if len(inlines) != len(panels):
            raise RuntimeError(f"Expected {len(panels)} placeholders, found {len(inlines)}.")
        relationships = {
            item.get("Id"): item
            for item in rels.getroot()
            if item.get("Type", "").endswith("/image")
        }

        for index, (inline, panel, placement) in enumerate(
            zip(inlines, panels, placements), start=1
        ):
            blip = inline.xpath(".//a:blip[@r:embed]", namespaces=NS)[0]
            rel_id = blip.get(f"{{{R}}}embed")
            relationship = relationships[rel_id]
            old_media = root / "word" / relationship.get("Target")
            extension = panel.source.suffix.lower()
            new_target = f"media/origin-panel-{index:03d}{extension}"
            new_media = root / "word" / new_target
            shutil.copyfile(panel.source, new_media)
            if old_media != new_media and old_media.exists():
                old_media.unlink()
            relationship.set("Target", new_target)
            ensure_content_type(content_types, extension)

            x0, y0, x1, y1 = panel.content_box
            crop = {
                "l": round(x0 / panel.width * 100000),
                "t": round(y0 / panel.height * 100000),
                "r": round((1 - x1 / panel.width) * 100000),
                "b": round((1 - y1 / panel.height) * 100000),
            }
            blip_fill = blip.getparent()
            src_rect = blip_fill.find(f"{{{A}}}srcRect")
            if src_rect is None:
                src_rect = etree.Element(f"{{{A}}}srcRect")
                blip_fill.insert(blip_fill.index(blip) + 1, src_rect)
            for key, value in crop.items():
                src_rect.set(key, str(value))

            width_emu = round(placement.width * scale * EMU_PER_INCH)
            height_emu = round(placement.height * scale * EMU_PER_INCH)
            extent = inline.find(f"{{{WP}}}extent")
            extent.set("cx", str(width_emu))
            extent.set("cy", str(height_emu))
            for inner_extent in inline.xpath(".//pic:spPr/a:xfrm/a:ext", namespaces=NS):
                inner_extent.set("cx", str(width_emu))
                inner_extent.set("cy", str(height_emu))
            doc_pr = inline.find(f"{{{WP}}}docPr")
            if doc_pr is not None:
                doc_pr.set("name", f"Origin panel {index}")
                doc_pr.set("descr", panel.source.name)
            x_emu = round((group_left + placement.anchor_x * scale) * EMU_PER_INCH)
            y_emu = round((group_top + placement.anchor_y * scale) * EMU_PER_INCH)
            inline_to_anchor(inline, x_emu, y_emu, index)

        document.write(str(document_path), xml_declaration=True, encoding="UTF-8", standalone="yes")
        rels.write(str(rels_path), xml_declaration=True, encoding="UTF-8", standalone="yes")
        content_types.write(
            str(content_types_path), xml_declaration=True, encoding="UTF-8", standalone="yes"
        )
        output.parent.mkdir(parents=True, exist_ok=True)
        rebuilt = output.with_suffix(output.suffix + ".building")
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as archive:
            for file in sorted(root.rglob("*")):
                if file.is_file():
                    archive.write(file, file.relative_to(root).as_posix())
        with zipfile.ZipFile(rebuilt) as archive:
            damaged = archive.testzip()
            if damaged:
                raise RuntimeError(f"Damaged DOCX member: {damaged}")
        rebuilt.replace(output)


def render_layout_preview(
    panels: list[Panel], placements: list[Placement], scale: float, output: Path
) -> None:
    dpi = 300
    width = math.ceil(max((item.anchor_x + item.width) * scale * dpi for item in placements))
    height = math.ceil(max((item.anchor_y + item.height) * scale * dpi for item in placements))
    canvas = Image.new("RGB", (width, height), "white")
    for panel, placement in zip(panels, placements):
        with Image.open(panel.preview) as image:
            cropped = image.convert("RGB").crop(panel.content_box)
        protected = Image.fromarray(panel.mask.astype(np.uint8) * 255, mode="L")
        target_width = max(1, round(placement.width * scale * dpi))
        target_height = max(1, round(placement.height * scale * dpi))
        cropped = cropped.resize((target_width, target_height), Image.Resampling.LANCZOS)
        protected = protected.resize((target_width, target_height), Image.Resampling.NEAREST)
        x = round(placement.anchor_x * scale * dpi)
        y = round(placement.anchor_y * scale * dpi)
        canvas.paste(cropped, (x, y), protected)
    canvas.save(output, "PNG")


def main() -> None:
    args = parse_args()
    if args.cols < 1:
        raise ValueError("--cols must be positive.")
    if args.safety_pt < 0:
        raise ValueError("--safety-pt cannot be negative.")
    sources = resolve_sources(args)
    rows = len(sources) // args.cols
    output = args.output.resolve()
    qa_dir = (args.qa_dir or output.with_name(output.stem + "_qa")).resolve()
    qa_dir.mkdir(parents=True, exist_ok=True)
    if args.preview_width < 800:
        raise ValueError("--preview-width must be at least 800 pixels for reliable text bounds.")
    previews = render_previews(sources, qa_dir, args.preview_width)
    overrides = load_overrides(args.frame_overrides_json)
    panels = [
        analyze_panel(source, preview, overrides.get(source.name))
        for source, preview in zip(sources, previews)
    ]
    edge_sources = [panel.source.name for panel in panels if panel.touches_source_edge]
    if edge_sources and not args.allow_edge_content:
        raise RuntimeError(
            "Visible content touches the source canvas edge and may already be clipped: "
            + ", ".join(edge_sources)
            + ". Re-export with more margin or pass --allow-edge-content after inspection."
        )
    placements, width_px, height_px, scale, clearance, frame_x, frame_y = choose_layout(
        panels,
        rows,
        args.cols,
        args.max_width_in,
        args.max_height_in,
        args.safety_pt,
    )
    group_width = width_px * scale
    group_height = height_px * scale
    text_width = args.page_width_in - args.left_margin_in - args.right_margin_in
    if group_width > text_width + 1e-6:
        raise RuntimeError(
            f"Requested figure width {group_width:.3f} in exceeds the text block {text_width:.3f} in."
        )
    group_left = args.left_margin_in + (text_width - group_width) / 2
    group_top = args.top_in

    with tempfile.TemporaryDirectory(prefix="origin_word_figure_base_") as temp:
        base = configure_document(args, len(panels), Path(temp))
        patch_docx(base, output, panels, placements, scale, group_left, group_top)

    preview_output = qa_dir / "layout-preview.png"
    render_layout_preview(panels, placements, scale, preview_output)
    report = {
        "output": str(output),
        "rows": rows,
        "columns": args.cols,
        "common_scale_in_per_preview_pixel": scale,
        "group_width_in": group_width,
        "group_height_in": group_height,
        "group_left_in": group_left,
        "group_top_in": group_top,
        "safety_pt_requested": args.safety_pt,
        "collision_clearance_px": clearance,
        "frame_origins_x_px": frame_x,
        "frame_origins_y_px": frame_y,
        "edge_content_warnings": edge_sources,
        "panels": [
            {
                "source": str(panel.source),
                "preview": str(panel.preview),
                "preview_size": [panel.width, panel.height],
                "content_box": list(panel.content_box),
                "plot_frame": list(panel.frame),
                "anchor_px": [placement.anchor_x, placement.anchor_y],
                "visible_size_px": [placement.width, placement.height],
            }
            for panel, placement in zip(panels, placements)
        ],
    }
    report_path = qa_dir / "layout-analysis.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output)
    print(preview_output)
    print(report_path)
    print(json.dumps({key: report[key] for key in report if key != "panels"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
